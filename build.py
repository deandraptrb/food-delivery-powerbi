"""
Data cleaning + star schema for the Online Food Ordering Dataset.

Input : data/online food delivery dataset.csv (raw from Kaggle)
Output: online_food_model.xlsx (FactResponse + Dim tables), ready to load into Power BI

Run from the repo folder:  python build.py
Requires: pandas, openpyxl
"""
from pathlib import Path
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
ROOT = Path(__file__).resolve().parent
RAW = ROOT / 'data' / 'online food delivery dataset.csv'
OUT = ROOT / 'online_food_model.xlsx'

# --- Cleaning ---
# The unnamed 14th column is identical to Output, so it is dropped
src=pd.read_csv(RAW).drop(columns='Unnamed: 13')
df=src.copy()
df['Feedback']=df['Feedback'].str.strip()
df['Occupation']=df['Occupation'].replace({'Self Employeed':'Self Employed','House wife':'Housewife'})
df['Monthly Income']=df['Monthly Income'].replace({'Below Rs.10000':'Below 10,000','10001 to 25000':'10,001 - 25,000','25001 to 50000':'25,001 - 50,000','More than 50000':'Above 50,000'})
# Exact duplicates are flagged, not deleted (Power BI measures exclude IsDuplicate = 1)
df.insert(0,'RespondentID',range(1,len(df)+1))
df['IsDuplicate']=df.drop(columns='RespondentID').duplicated().map({True:1,False:0})
bins=[0,20,23,26,29,40]; labels=['18-20','21-23','24-26','27-29','30+']
df['AgeGroup']=pd.cut(df.Age,bins=bins,labels=labels).astype(str)

# --- Dimension tables (Key + label + SortOrder) ---
def dim(col,order,key):
    d=pd.DataFrame({col:order}); d.insert(0,key,range(1,len(d)+1)); d['SortOrder']=d[key]; return d
dims={
 'DimOccupation':dim('Occupation',['Student','Employee','Self Employed','Housewife'],'OccupationKey'),
 'DimIncome':dim('Monthly Income',['No Income','Below 10,000','10,001 - 25,000','25,001 - 50,000','Above 50,000'],'IncomeKey'),
 'DimEducation':dim('Education',['Uneducated','School','Graduate','Post Graduate','Ph.D'],'EducationKey'),
 'DimMaritalStatus':dim('Marital Status',['Single','Married','Prefer not to say'],'MaritalKey'),
 'DimGender':dim('Gender',['Female','Male'],'GenderKey'),
 'DimAgeGroup':dim('Age Group',labels,'AgeGroupKey'),
}
df=df.rename(columns={'Educational Qualifications':'Education','AgeGroup':'Age Group'})
for name,d in dims.items():
    col,key=d.columns[1],d.columns[0]
    assert set(df[col])<=set(d[col]), (name,set(df[col])-set(d[col]))
    df=df.merge(d[[key,col]],on=col,how='left')
loc=df.groupby('Pin code',as_index=False).agg(Latitude=('latitude','mean'),Longitude=('longitude','mean'))
loc['City']='Bengaluru'; loc['Country']='India'
fact=df[['RespondentID','Age','AgeGroupKey','GenderKey','MaritalKey','OccupationKey','IncomeKey','EducationKey','Family size','Pin code','latitude','longitude','Customer Type','Output','Feedback','IsDuplicate']].rename(columns={'Family size':'FamilySize','Pin code':'PinCode','latitude':'Latitude','longitude':'Longitude','Customer Type':'CustomerType','Output':'OrderAgain'})
fact=fact.sort_values('RespondentID')
loc=loc.rename(columns={'Pin code':'PinCode'})
notes=pd.DataFrame({'Item':[
 'Source','Grain','Cleaning: Feedback','Cleaning: last column','Cleaning: labels','Duplicates','CustomerType warning','Relationships','Sort order'],
 'Detail':[
 'Kaggle: srisyra02/online-food-ordering-dataset (388 rows, survey of online food ordering customers in Bengaluru)',
 'FactResponse = 1 row per survey response',
 'Trailing space removed from "Negative "',
 'Unnamed 14th column dropped (identical to Output in all 388 rows)',
 '"Self Employeed" -> "Self Employed", "House wife" -> "Housewife", income bands relabeled',
 '103 rows are exact copies of an earlier row (flag IsDuplicate = 1). Measures exclude them by default; 285 unique responses remain',
 'CustomerType is fully determined by family size (1 = New, 2-3 = Regular, 4-6 = Frequent), so it is not real ordering behavior. Do not present it as order frequency',
 'Each Dim*Key -> FactResponse (1:*, single direction). DimLocation[PinCode] -> FactResponse[PinCode]',
 'In Power BI set each label column "Sort by column" = SortOrder']})
# --- Export ---
path=OUT
with pd.ExcelWriter(path,engine='openpyxl') as w:
    fact.to_excel(w,sheet_name='FactResponse',index=False)
    for n,d in dims.items(): d.to_excel(w,sheet_name=n,index=False)
    loc.to_excel(w,sheet_name='DimLocation',index=False)
    notes.to_excel(w,sheet_name='README',index=False)
wb=load_workbook(path)
for ws in wb:
    for row in ws.iter_rows():
        for c in row: c.font=Font(name='Arial',size=10,bold=(c.row==1))
    for c in ws[1]: c.fill=PatternFill('solid',fgColor='EAF5F5')
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width=min(max(len(str(c.value or '')) for c in col)+2,110)
    ws.freeze_panes='A2'
wb.save(path)
print(fact.head().to_string()); print(len(fact), fact.IsDuplicate.sum(), fact.isna().sum().sum(), len(loc))
